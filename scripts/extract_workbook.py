from __future__ import annotations

import json
import sys
from pathlib import Path


APP_DIR = Path(__file__).resolve().parents[1]
PROJECT_DIR = APP_DIR.parent
sys.path.insert(0, str(PROJECT_DIR))

from scoring_core import (  # noqa: E402
    CATEGORY_NAMES,
    DIAGNOSIS_SHEETS,
    _matching_options,
    clean,
    load_questions,
    normalized,
)
from openpyxl import load_workbook  # noqa: E402


WORKBOOK_PATH = PROJECT_DIR / "data" / "OTO Dizziness Questionnaire -Nebula cloud 5-9-241.xlsx"
OUTPUT_PATH = APP_DIR / "dist" / "questionnaire-data.js"


def main() -> None:
    questions = load_questions(WORKBOOK_PATH)
    questions_by_id = {question.question_id: question for question in questions}
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)

    rules: list[dict[str, object]] = []
    referenced_ids: set[str] = set()

    for sheet_name in DIAGNOSIS_SHEETS:
        category = CATEGORY_NAMES.get(sheet_name, sheet_name.strip())
        eligible_rows: list[tuple[object, list[str]]] = []

        for row in workbook[sheet_name].iter_rows(values_only=True):
            question_id = clean(row[0] if row else "")
            question = questions_by_id.get(question_id)
            if not question:
                continue

            answers = [
                answer
                for answer in _matching_options(row[2] if len(row) > 2 else "", question)
                if normalized(answer) != "no"
            ]
            if answers:
                eligible_rows.append((question, answers))

        row_weight = 100.0 / max(len(eligible_rows), 1)
        for question, answers in eligible_rows:
            answer_weight = row_weight / len(answers)
            referenced_ids.add(question.question_id)
            for answer in answers:
                rules.append(
                    {
                        "category": category,
                        "questionId": question.question_id,
                        "answer": answer,
                        "weight": round(answer_weight, 3),
                        "editable": question.kind == "single",
                        "source": f"Workbook sheet: {sheet_name.strip()}",
                    }
                )

    workbook.close()

    limited_questions = [
        {
            "id": question.question_id,
            "prompt": question.prompt,
            "kind": question.kind,
            "options": question.options,
        }
        for question in questions
        if question.question_id in referenced_ids
    ]

    payload = {
        "version": 1,
        "sourceWorkbook": WORKBOOK_PATH.name,
        "questions": limited_questions,
        "rules": rules,
        "categories": sorted({rule["category"] for rule in rules}, key=str.casefold),
        "policy": {
            "defaultWeightsOnly": True,
            "machineLearningWeightsIncluded": False,
            "noAnswerWeightsExcluded": True,
            "multiSelectWeightsEditable": False,
        },
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(
        "window.NEBULA_DATA = " + json.dumps(payload, indent=2, ensure_ascii=False) + ";\n",
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "questions": len(limited_questions),
                "rules": len(rules),
                "editableRules": sum(bool(rule["editable"]) for rule in rules),
                "fixedRules": sum(not bool(rule["editable"]) for rule in rules),
                "noAnswerRules": sum(normalized(rule["answer"]) == "no" for rule in rules),
                "output": str(OUTPUT_PATH),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
